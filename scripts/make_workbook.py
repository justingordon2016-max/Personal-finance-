#!/usr/bin/env python3
"""Build reports/budget.xlsx from data/budget.csv.

One editable sheet per block (Fixed, Variable, Savings) feeding a Summary that
recalculates. Edit the blue amounts and everything downstream follows.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SALARY = 70000.00

FONT = "Arial"
INK = Font(name=FONT, size=10)
INPUT = Font(name=FONT, size=10, color="0000FF")          # editable
FORMULA = Font(name=FONT, size=10)                         # calculated
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
GROUP = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="595959")

HEADFILL = PatternFill("solid", fgColor="0A6E62")
GROUPFILL = PatternFill("solid", fgColor="E4EFED")
TOTALFILL = PatternFill("solid", fgColor="D5E5E2")
FLAGFILL = PatternFill("solid", fgColor="FBF0DA")

MONEY = '#,##0.00;(#,##0.00);-'
PCT = '0.0%'
THIN = Side(style="thin", color="C8D2D0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def load():
    rows = list(csv.DictReader((ROOT / "data" / "budget.csv").open()))
    for r in rows:
        r["amount"] = float(r["amount"])
    return rows


def sheet(wb, rows, kind, title, cols):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = ("Edit the blue amounts. Annual and the Summary tab recalculate. "
                "Amounts are per month.")
    ws["A2"].font = NOTE
    ws.append([])

    hdr = 4
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=hdr, column=i, value=c)
        cell.font, cell.fill, cell.border = HEAD, HEADFILL, BOX
        cell.alignment = Alignment(horizontal="right" if i >= 5 else "left")

    r = hdr + 1
    first_data = r
    cats, group_rows = [], []
    for cat in dict.fromkeys(x["category"] for x in rows if x["type"] == kind):
        items = sorted([x for x in rows if x["type"] == kind and x["category"] == cat],
                       key=lambda x: -x["amount"])
        gr = r
        ws.cell(row=r, column=1, value=cat.upper())
        for i in range(1, len(cols) + 1):
            ws.cell(row=r, column=i).fill = GROUPFILL
            ws.cell(row=r, column=i).font = GROUP
            ws.cell(row=r, column=i).border = BOX
        start = r + 1
        r += 1
        for it in items:
            ws.cell(row=r, column=1, value=it["item"]).font = INK
            ws.cell(row=r, column=2, value=it["frequency"]).font = INK
            ws.cell(row=r, column=3, value=it["day"]).font = INK
            ws.cell(row=r, column=4, value=it["account"]).font = INK
            a = ws.cell(row=r, column=5, value=it["amount"])
            a.font, a.number_format = INPUT, MONEY
            y = ws.cell(row=r, column=6, value=f"=E{r}*12")
            y.font, y.number_format = FORMULA, MONEY
            p = ws.cell(row=r, column=7, value=f"=E{r}/Summary!$B$4")
            p.font, p.number_format = FORMULA, PCT
            n = ws.cell(row=r, column=8,
                        value=(it["status"].upper() + " — " if it["status"] else "")
                        + it["basis"])
            n.font = NOTE
            if it["status"]:
                for i in range(1, len(cols) + 1):
                    ws.cell(row=r, column=i).fill = FLAGFILL
            for i in range(1, len(cols) + 1):
                ws.cell(row=r, column=i).border = BOX
            r += 1
        ws.cell(row=gr, column=5, value=f"=SUM(E{start}:E{r-1})").number_format = MONEY
        ws.cell(row=gr, column=5).font = GROUP
        ws.cell(row=gr, column=6, value=f"=E{gr}*12").number_format = MONEY
        ws.cell(row=gr, column=6).font = GROUP
        ws.cell(row=gr, column=7, value=f"=E{gr}/Summary!$B$4").number_format = PCT
        ws.cell(row=gr, column=7).font = GROUP
        cats.append(cat)
        group_rows.append(gr)

    ws.cell(row=r, column=1, value=f"TOTAL {title.upper()}").font = GROUP
    tot = ws.cell(row=r, column=5, value="+".join(f"E{g}" for g in group_rows))
    tot.value = "=" + "+".join(f"E{g}" for g in group_rows)
    tot.font, tot.number_format = GROUP, MONEY
    ws.cell(row=r, column=6, value=f"=E{r}*12").number_format = MONEY
    ws.cell(row=r, column=6).font = GROUP
    ws.cell(row=r, column=7, value=f"=E{r}/Summary!$B$4").number_format = PCT
    ws.cell(row=r, column=7).font = GROUP
    for i in range(1, len(cols) + 1):
        ws.cell(row=r, column=i).fill = TOTALFILL
        ws.cell(row=r, column=i).border = BOX

    for col, w in zip("ABCDEFGH", (40, 12, 12, 12, 14, 14, 11, 58)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{hdr+1}"
    return ws, r


def main():
    rows = load()
    wb = Workbook()
    wb.remove(wb.active)

    summary = wb.create_sheet("Summary")
    cols = ["Item", "Frequency", "Clears", "Account", "Per month", "Per year",
            "% of salary", "Basis / note"]
    _, fixed_total = sheet(wb, rows, "Fixed", "Fixed", cols)
    _, var_total = sheet(wb, rows, "Variable", "Variable", cols)
    _, sav_total = sheet(wb, rows, "Savings", "Savings", cols)

    ws = summary
    ws["A1"] = "Budget — September 2026"
    ws["A1"].font = TITLE
    ws["A2"] = ("Blue cells are inputs. Everything else is a formula. "
                "Edit amounts on the Fixed, Variable and Savings tabs.")
    ws["A2"].font = NOTE

    ws["A4"] = "Salary (net, monthly)"
    ws["A4"].font = GROUP
    ws["B4"] = SALARY
    ws["B4"].font, ws["B4"].number_format = INPUT, MONEY
    ws["C4"] = "Source: FNB statements 135-138, constant at R70,000"
    ws["C4"].font = NOTE

    plan = [
        ("Less: Fixed expenses",   f"=-Fixed!E{fixed_total}",    None),
        ("LEFT AFTER FIXED",       "=B4+B6",                     "sub"),
        ("Less: Variable expenses", f"=-Variable!E{var_total}",  None),
        ("LEFT AFTER FIXED + VARIABLE", "=B7+B8",                "sub"),
        ("Less: Savings instructions", f"=-Savings!E{sav_total}", None),
        ("MONTHLY GAP",            "=B9+B10",                    "end"),
    ]
    r = 6
    for label, formula, kind in plan:
        ws.cell(row=r, column=1, value=label).font = GROUP if kind else INK
        c = ws.cell(row=r, column=2, value=formula)
        c.font, c.number_format = (GROUP if kind else FORMULA), MONEY
        pc = ws.cell(row=r, column=3, value=f"=B{r}/$B$4")
        pc.font, pc.number_format = NOTE, PCT
        if kind == "end":
            for i in (1, 2, 3):
                ws.cell(row=r, column=i).fill = TOTALFILL
        elif kind == "sub":
            for i in (1, 2, 3):
                ws.cell(row=r, column=i).fill = GROUPFILL
        for i in (1, 2, 3):
            ws.cell(row=r, column=i).border = BOX
        r += 1

    notes = [
        "",
        "How the three blocks are defined",
        "Fixed — a standing instruction at a rate you cannot change this month, shown at its "
        "most recent amount rather than an average.",
        "Variable — you choose when and how much, so these are 4-cycle averages. Prepaid "
        "electricity and airtime sit here for that reason.",
        "Savings — standing instructions too, but money kept rather than spent, so they are "
        "kept in their own block.",
        "",
        "Open items",
        "Payments to a second credit account (R7,750/mo) are an average of five irregular and "
        "accelerating payments. No statement for that account exists yet.",
        "Ril Res001 (R850/mo) first appeared on 1 September and is assumed recurring.",
        "A R200,000 capital repayment on 9 September may cut the structured-loan instalment "
        "from 30 September. Watch that debit.",
        "EasyEquities TFSA at R3,006.22/mo is R74.64/yr over the R36,000 annual limit; excess "
        "contributions are taxed at 40%.",
        "",
        "Excluded: R26,060.66 of out-of-pocket medical in June — a one-off event, not a "
        "monthly cost.",
        "Source: FNB cheque statements 135-138 and card statements 130-132, all reconciled to "
        "the printed balances.",
    ]
    for i, n in enumerate(notes):
        c = ws.cell(row=14 + i, column=1, value=n)
        c.font = GROUP if n in ("How the three blocks are defined", "Open items") else NOTE
    for col, w in zip("ABC", (46, 16, 13)):
        ws.column_dimensions[col].width = w

    out = ROOT / "reports" / "budget.xlsx"
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
